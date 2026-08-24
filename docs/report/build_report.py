"""Build the editable Assignment 01 Microsoft Word technical report.

The report body is sourced from TECHNICAL_REPORT.md. Existing notebook figures are
embedded as images, while all report text, headings, captions, and tables remain
native editable Word content.
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = REPORT_DIR.parent.parent
MARKDOWN_PATH = REPORT_DIR / "TECHNICAL_REPORT.md"
DOCX_PATH = REPORT_DIR / "TECHNICAL_REPORT.docx"
ASSET_DIR = REPORT_DIR / "assets"

INK = "20313A"
TEAL = "176B6A"
BLUE = "3E6C8A"
ORANGE = "D9822B"
PALE = "F3F7F6"
LINE = "AAB8BD"


def draw_box(ax, x, y, w, h, text, *, face=PALE, edge=TEAL, fontsize=10):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.02,rounding_size=0.025",
        linewidth=1.5,
        edgecolor=f"#{edge}",
        facecolor=f"#{face}",
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fontsize, color=f"#{INK}")
    return patch


def arrow(ax, start, end, *, color=BLUE, connectionstyle="arc3,rad=0"):
    ax.add_patch(FancyArrowPatch(
        start, end,
        arrowstyle="-|>",
        mutation_scale=13,
        linewidth=1.4,
        color=f"#{color}",
        connectionstyle=connectionstyle,
    ))


def generate_system_architecture() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "system_architecture.png"
    fig, ax = plt.subplots(figsize=(14, 8.2))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.965, "Deployed Intelligent Application Architecture", ha="center", va="top", fontsize=18, weight="bold", color=f"#{INK}")

    draw_box(ax, 0.07, 0.72, 0.22, 0.13, "React Web\nVercel", face="EAF3F6", edge=BLUE, fontsize=12)
    draw_box(ax, 0.07, 0.39, 0.22, 0.13, "Expo React Native\nExpo Go", face="EAF3F6", edge=BLUE, fontsize=12)
    draw_box(ax, 0.40, 0.56, 0.22, 0.16, "FastAPI API\nRender", face="EAF6F3", edge=TEAL, fontsize=13)
    ax.text(0.34, 0.78, "HTTPS", ha="center", color=f"#{BLUE}", fontsize=10, weight="bold")
    ax.text(0.34, 0.48, "HTTPS", ha="center", color=f"#{BLUE}", fontsize=10, weight="bold")
    arrow(ax, (0.29, 0.785), (0.40, 0.665))
    arrow(ax, (0.29, 0.455), (0.40, 0.61))

    draw_box(ax, 0.72, 0.72, 0.22, 0.13, "Diabetes Pipeline\nRF Classifier", face="FFF4E7", edge=ORANGE, fontsize=11)
    draw_box(ax, 0.72, 0.50, 0.22, 0.13, "House Pipeline\nRF Regressor", face="FFF4E7", edge=ORANGE, fontsize=11)
    draw_box(ax, 0.72, 0.27, 0.22, 0.13, "Neo4j AuraDB\nDiabetes KG", face="F2ECF8", edge="71558F", fontsize=11)
    arrow(ax, (0.62, 0.66), (0.72, 0.785))
    arrow(ax, (0.62, 0.63), (0.72, 0.565))
    arrow(ax, (0.62, 0.59), (0.72, 0.335))

    draw_box(ax, 0.40, 0.18, 0.22, 0.12, "Prediction Responses\nJSON over HTTPS", face="F7F7F4", edge="758187", fontsize=11)
    arrow(ax, (0.83, 0.72), (0.62, 0.30), connectionstyle="arc3,rad=-0.22")
    arrow(ax, (0.83, 0.50), (0.62, 0.27), connectionstyle="arc3,rad=-0.12")
    arrow(ax, (0.40, 0.24), (0.29, 0.72), connectionstyle="arc3,rad=0.20")
    arrow(ax, (0.40, 0.22), (0.29, 0.45), connectionstyle="arc3,rad=0.10")
    ax.text(0.51, 0.08, "Saved fitted scikit-learn Pipelines perform all learned preprocessing and inference.", ha="center", fontsize=10, color=f"#{INK}")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def generate_ml_pipeline() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "ml_pipeline.png"
    fig, ax = plt.subplots(figsize=(14, 9))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.27, 0.96, "TRAINING", ha="center", va="top", fontsize=17, weight="bold", color=f"#{TEAL}")
    ax.text(0.75, 0.96, "INFERENCE", ha="center", va="top", fontsize=17, weight="bold", color=f"#{BLUE}")

    training = [
        "Dataset", "Raw Feature\nRepresentation", "Train/Test Split",
        "Preprocessing\ninside Pipeline", "Multiple Traditional\nML Models",
        "Controlled\nExperiments", "Final Selection",
        "Saved sklearn\nPipeline",
    ]
    ys = [0.84, 0.74, 0.64, 0.54, 0.44, 0.34, 0.24, 0.12]
    for index, (label, y) in enumerate(zip(training, ys)):
        draw_box(ax, 0.12, y, 0.30, 0.075, label, face="EEF7F4", edge=TEAL, fontsize=10.5)
        if index:
            arrow(ax, (0.27, ys[index - 1]), (0.27, y + 0.075), color=TEAL)

    inference = ["Raw User Input", "Same Saved Pipeline", "Prediction", "Web / Mobile"]
    iys = [0.72, 0.54, 0.36, 0.18]
    for index, (label, y) in enumerate(zip(inference, iys)):
        draw_box(ax, 0.61, y, 0.28, 0.095, label, face="EEF4F8", edge=BLUE, fontsize=11)
        if index:
            arrow(ax, (0.75, iys[index - 1]), (0.75, y + 0.095), color=BLUE)

    arrow(ax, (0.42, 0.157), (0.61, 0.587), color=ORANGE, connectionstyle="arc3,rad=-0.18")
    ax.text(0.51, 0.42, "load once", ha="center", va="center", fontsize=10, weight="bold", color=f"#{ORANGE}", rotation=28)
    ax.text(0.51, 0.045, "The same learned imputation, scaling, encoding, and estimator are preserved across training and inference.", ha="center", fontsize=10, color=f"#{INK}")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_node, end])


def add_hyperlink(paragraph, label: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = label
    run.extend([props, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^)]+\)|<https?://[^>]+>)")


def add_inline(paragraph, text: str, *, bold_all: bool = False) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold = bold_all
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token[1:-1]
            add_hyperlink(paragraph, url, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = bold_all


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.7)

    heading_specs = {
        "Heading 1": (17, TEAL, 12, 6),
        "Heading 2": (15, INK, 10, 4),
        "Heading 3": (13, INK, 8, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.first_line_indent = Cm(0)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.left_indent = Cm(0.5)
    code_style.paragraph_format.right_indent = Cm(0.3)
    code_style.paragraph_format.space_after = Pt(5)
    code_style.paragraph_format.first_line_indent = Cm(0)


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def add_header_footer(section, *, first_page=False) -> None:
    if first_page:
        section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Intelligent System Development – Assignment 01")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("66747A")

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    add_field(paragraph, "PAGE", "1")


def add_title_page(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INTELLIGENT SYSTEM DEVELOPMENT")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.paragraph_format.space_after = Pt(10)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ASSIGNMENT 01")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("From Data Representation to Intelligent Applications")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)

    for _ in range(2):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "Systems:\n1. Diabetes Classification System\n2. Vietnam House Price Prediction System")
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
    p.paragraph_format.line_spacing = 1.5

    for _ in range(3):
        document.add_paragraph()
    for text in [
        "Student: Nguyễn Thành Trung",
        "Student ID: B23DCCN861",
        "Class: D23CTPM01-B",
        "Lecturer: _________________________________",
    ]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("August 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)


def add_toc(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("TABLE OF CONTENTS")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    toc_p = document.add_paragraph()
    toc_p.paragraph_format.first_line_indent = Cm(0)
    add_field(toc_p, 'TOC \\o "1-3" \\h \\z \\u', "[Right-click and update the Table of Contents in Microsoft Word]")
    note = document.add_paragraph("If page numbers are not visible, right-click the field above and select Update Field → Update entire table.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = Cm(0)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(10)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value, bold_all=(row_index == 0))
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9 if len(rows[0]) >= 5 else 9.5)
            if row_index == 0:
                set_cell_shading(cell, "DDECE9")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, source: str, caption: str) -> None:
    path = (REPORT_DIR / source).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Figure not found: {path}")
    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = height_px / width_px
    max_width = 6.25
    max_height = 7.25
    width = min(max_width, max_height / ratio)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(style="Caption")
    add_inline(caption_paragraph, caption)


def parse_markdown(document: Document, markdown: str) -> tuple[int, int]:
    body = markdown.split("<!-- REPORT BODY -->", 1)[1]
    lines = body.splitlines()
    index = 0
    tables = 0
    figures = 0
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []
    major_breaks = {
        "4. Dataset", "8. Traditional Machine Learning Models", "10. Experimental Results",
        "12. Intelligent Application Architecture", "16. Diabetes Knowledge Graph",
        "18. System Demonstration", "19. Limitations", "22. Reproducibility",
        "23. References", "Appendix A – API Endpoints",
    }

    def flush_paragraph():
        nonlocal paragraph_lines
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines).strip()
        if text:
            style = "Caption" if re.match(r"^\*\*(Table|Figure) \d+\.", text) else None
            paragraph = document.add_paragraph(style=style)
            if style:
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, text)
        paragraph_lines = []

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = document.add_paragraph("\n".join(code_lines), style="Code Block")
                paragraph.paragraph_format.keep_together = True
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped == "<!-- PAGE BREAK -->":
            flush_paragraph()
            document.add_page_break()
            index += 1
            continue
        image_match = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            flush_paragraph()
            add_figure(document, image_match.group(2), image_match.group(1))
            figures += 1
            index += 1
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            title = heading_match.group(2)
            if level == 1 and title in major_breaks and document.paragraphs:
                document.add_page_break()
            document.add_heading(title, level=level)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1].strip()):
            flush_paragraph()
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip().replace("<br>", "\n") for cell in row.strip("|").split("|")]
                for row in table_lines
            ]
            add_markdown_table(document, rows)
            tables += 1
            continue
        list_match = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if list_match or numbered_match:
            flush_paragraph()
            style = "List Bullet" if list_match else "List Number"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, (list_match or numbered_match).group(1))
            index += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.right_indent = Cm(0.5)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, stripped[2:])
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()
    return tables, figures


def build() -> tuple[int, int]:
    generate_system_architecture()
    generate_ml_pipeline()
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    configure_section(section)
    add_header_footer(section, first_page=True)
    document.core_properties.title = "Intelligent System Development – Assignment 01 Technical Report"
    document.core_properties.subject = "Diabetes classification and Vietnam house price regression intelligent systems"
    document.core_properties.author = "Student: ______________________________"
    document.core_properties.keywords = "intelligent systems, machine learning, FastAPI, React, Expo, Neo4j"

    add_title_page(document)
    document.add_page_break()
    add_toc(document)
    document.add_page_break()
    markdown = MARKDOWN_PATH.read_text(encoding="utf-8")
    tables, figures = parse_markdown(document, markdown)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(DOCX_PATH)
    return tables, figures


def validate(tables: int, figures: int) -> None:
    if not zipfile.is_zipfile(DOCX_PATH):
        raise RuntimeError("Generated file is not a valid DOCX ZIP package")
    document = Document(DOCX_PATH)
    headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
    heading_text = {p.text.strip() for p in headings}
    required = {
        "1. Introduction", "10. Experimental Results", "17. Deployment Architecture",
        "23. References", "Appendix C – Demo Input Cases",
    }
    missing = required - heading_text
    if missing:
        raise RuntimeError(f"Missing required headings: {sorted(missing)}")
    relationships = document.part.rels.values()
    image_relationships = [rel for rel in relationships if "image" in rel.reltype]
    body_text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    full_text = body_text + "\n" + table_text
    if "TOC \\o" not in document._element.xml:
        raise RuntimeError("TOC field was not created")
    if "PAGE" not in document.sections[0].footer._element.xml:
        raise RuntimeError("Page-number field was not created")
    if len(document.paragraphs) < 150 or not document.tables or not headings or not image_relationships:
        raise RuntimeError("DOCX structural validation failed")
    if "R² = 0.4738 does not mean 47.38% accuracy" not in body_text:
        raise RuntimeError("Regression accuracy warning is missing")
    critical_facts = [
        "Student: Nguyễn Thành Trung", "Student ID: B23DCCN861", "Class: D23CTPM01-B",
        "Lecturer: _________________________________",
        "768 observations", "500 class-0 and 268 class-1", "29.56%", "48.70%",
        "0.7013", "0.7338", "0.6818", "0.7597", "0.6408",
        "0.6462", "0.6243", "0.6061", "[[85, 15], [24, 30]]",
        "30,229 rows", "1.8438", "4.8760", "2.2082", "44.43%",
        "1.4782", "1.3627", "1.3976", "1.3009", "1.2868",
        "1.6565", "1.6526", "1.6078", "1.2529", "2.5659", "1.6018",
        "0.4738", "27.36%", "17 nodes and 17 relationships",
        "Responsive breakpoints adapt navigation", "Its five screens are Home, Diabetes, House Price, Knowledge Graph, and About",
        "touch-oriented node explorer",
    ]
    missing_facts = [fact for fact in critical_facts if fact not in full_text]
    if missing_facts:
        raise RuntimeError(f"Critical source facts missing from DOCX: {missing_facts}")
    first_section = document.sections[0]
    if abs(first_section.page_width.cm - 21.0) > 0.05 or abs(first_section.page_height.cm - 29.7) > 0.05:
        raise RuntimeError("Page size is not A4")
    margins = [
        first_section.top_margin.cm, first_section.bottom_margin.cm,
        first_section.left_margin.cm, first_section.right_margin.cm,
    ]
    if any(abs(margin - 2.0) > 0.05 for margin in margins):
        raise RuntimeError("Page margins are not approximately 2 cm")
    words = len(re.findall(r"\b\w+[\w²-]*\b", body_text, flags=re.UNICODE))
    estimated_pages = round(2 + words / 510 + figures * 0.48 + tables * 0.16)
    print(f"DOCX: {DOCX_PATH}")
    print(f"paragraphs={len(document.paragraphs)} headings={len(headings)} tables={len(document.tables)}")
    print(f"figures={len(image_relationships)} words~={words} estimated_pages~={estimated_pages}")
    print("valid_docx_zip=True toc_field=True page_field=True required_sections=True critical_facts=True a4=True")


if __name__ == "__main__":
    if "--validate-only" in sys.argv:
        existing = Document(DOCX_PATH)
        image_count = len([rel for rel in existing.part.rels.values() if "image" in rel.reltype])
        validate(len(existing.tables), image_count)
    else:
        table_count, figure_count = build()
        validate(table_count, figure_count)
