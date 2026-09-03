"""Build and validate the editable Vietnamese Assignment 02 report."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image, ImageOps
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parent
ROOT = REPORT_DIR.parents[1]
MARKDOWN = REPORT_DIR / "ASSIGNMENT_02_REPORT.md"
DOCX = REPORT_DIR / "ASSIGNMENT_02_REPORT.docx"
ASSETS = REPORT_DIR / "assets"

INK = "18312C"
TEAL = "1D5A54"
ORANGE = "D96F32"
MUTED = "65736F"
PALE = "EEF4F1"

TABLE_CAPTIONS = [
    "Bảng 1. So sánh tổng quan ba hệ thống thông minh.",
    "Bảng 2. Tóm tắt raw form và numerical representation.",
    "Bảng 3. Controlled representation experiment của Ecommerce.",
    "Bảng 4. Validation comparison của sáu Ecommerce models.",
    "Bảng 5. Final test metrics và phạm vi diễn giải.",
    "Bảng 6. Diabetes baseline và năm model trên holdout mô tả.",
    "Bảng 7. Hai benchmark House bổ sung trên training folds.",
    "Bảng 8. Ecommerce representation evidence trong phụ lục.",
    "Bảng 9. Ecommerce six-model evidence trong phụ lục.",
]


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def generate_architecture() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / "system_architecture.png"
    fig, ax = plt.subplots(figsize=(13, 6.4))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    def box(x, y, w, h, text, color=TEAL, face=PALE, size=11):
        patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.025",
                               linewidth=1.6, edgecolor=f"#{color}", facecolor=f"#{face}")
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, color=f"#{INK}")

    def arrow(start, end):
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14,
                                     linewidth=1.5, color=f"#{MUTED}"))

    ax.text(0.5, 0.95, "KIẾN TRÚC TRIỂN KHAI ASSIGNMENT 02", ha="center", va="top",
            fontsize=18, weight="bold", color=f"#{INK}")
    box(0.04, 0.62, 0.18, 0.16, "Web React\nVercel", color="3E6C8A", face="EAF3F6")
    box(0.04, 0.27, 0.18, 0.16, "Mobile Expo\nExpo Go", color="3E6C8A", face="EAF3F6")
    box(0.34, 0.43, 0.22, 0.22, "FastAPI · Render\nValidation\nRaw DataFrame", size=12)
    arrow((0.22, 0.70), (0.34, 0.58))
    arrow((0.22, 0.35), (0.34, 0.49))
    for y, text in [(0.68, "Diabetes Pipeline\nB × 6"), (0.43, "House Pipeline\nB × 83"), (0.18, "Ecommerce Pipeline\nB × 12.005")]:
        box(0.70, y, 0.24, 0.15, text, color=ORANGE, face="FFF4EA")
    arrow((0.56, 0.57), (0.70, 0.755))
    arrow((0.56, 0.54), (0.70, 0.505))
    arrow((0.56, 0.50), (0.70, 0.255))
    ax.text(0.50, 0.08, "User input → schema validation → saved fitted Pipeline → prediction → result",
            ha="center", fontsize=11, color=f"#{MUTED}")
    fig.tight_layout()
    fig.savefig(path, dpi=190, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_contact_sheet(paths: list[Path], target: Path, columns: int, width: int) -> None:
    images = []
    for path in paths:
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
            height = round(image.height * width / image.width)
            images.append(image.resize((width, height), Image.Resampling.LANCZOS))
    rows = (len(images) + columns - 1) // columns
    gap = 24
    row_heights = []
    for row in range(rows):
        row_heights.append(max(image.height for image in images[row * columns:(row + 1) * columns]))
    canvas = Image.new("RGB", (columns * width + (columns - 1) * gap,
                               sum(row_heights) + (rows - 1) * gap), "white")
    y = 0
    for row in range(rows):
        x = 0
        for image in images[row * columns:(row + 1) * columns]:
            canvas.paste(image, (x, y))
            x += width + gap
        y += row_heights[row] + gap
    target.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(target, quality=94)


def generate_assets() -> None:
    generate_architecture()
    make_contact_sheet([
        ROOT / "figures/png/web/Home.png",
        ROOT / "figures/png/web/Diabetes.png",
        ROOT / "figures/png/web/HousePricing.png",
        ROOT / "figures/png/web/Ecommerce.png",
    ], ASSETS / "web_production.png", columns=2, width=960)
    make_contact_sheet([
        ROOT / "figures/png/mobile/Diabetes/diabetes2.png",
        ROOT / "figures/png/mobile/HousePricing/house4.png",
        ROOT / "figures/png/mobile/Ecommerce/e2.png",
    ], ASSETS / "mobile_results.png", columns=3, width=420)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.7)

    for name, size, color in (("Heading 1", 16, TEAL), ("Heading 2", 14, INK), ("Heading 3", 12.5, INK)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_after = Pt(5)

    for style_name, fill, font_size in (
        ("Code Block", "F2F6F5", 9.0),
        ("Output Block", "EEF5F9", 9.0),
    ):
        if style_name not in document.styles:
            block_style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            block_style = document.styles[style_name]
        block_style.font.name = "Consolas"
        block_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        block_style.font.size = Pt(font_size)
        block_style.paragraph_format.left_indent = Cm(0.55)
        block_style.paragraph_format.right_indent = Cm(0.35)
        block_style.paragraph_format.first_line_indent = Cm(0)
        block_style.paragraph_format.line_spacing = 1.0
        block_style.paragraph_format.space_before = Pt(3)
        block_style.paragraph_format.space_after = Pt(6)
        properties = block_style._element.get_or_add_pPr()
        shading = properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            properties.append(shading)
        shading.set(qn("w:fill"), fill)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("Intelligent System Development · Assignment 02")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Trang ")
    add_field(footer, "PAGE", "1")


def add_cover(document: Document) -> None:
    for _ in range(2):
        document.add_paragraph()
    for text, size, color in [
        ("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 16, INK),
        ("INTELLIGENT SYSTEM DEVELOPMENT", 20, TEAL),
        ("BÀI TẬP 02", 20, INK),
        ("FROM DATA REPRESENTATION TO A DEPLOYABLE INTELLIGENT SYSTEM", 15, ORANGE),
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
    for _ in range(3):
        document.add_paragraph()
    for text in [
        "Sinh viên: Nguyễn Thành Trung",
        "Mã sinh viên: B23DCCN861",
        "Lớp: D23CTPM01-B",
        "Giảng viên: Dinh Que Tran, Ph.D., Assoc. Prof.",
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(22)
    paragraph.add_run("2026").font.size = Pt(13)


INLINE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*.+?\*\*|`.+?`)")


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_table(document: Document, rows: list[list[str]], caption_text: str) -> None:
    caption = document.add_paragraph(style="Caption")
    caption.paragraph_format.keep_with_next = True
    add_inline(caption, caption_text)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_header(table.rows[0])
    font_size = 7.5 if len(rows[0]) >= 6 else 8.5 if len(rows[0]) >= 4 else 9.5
    for row_index, values in enumerate(rows):
        prevent_row_split(table.rows[row_index])
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                shade(cell, "DDECE7")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, source: str, caption: str) -> None:
    path = (REPORT_DIR / source).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        ratio = image.height / image.width
    max_width = 6.25
    max_height = 4.2 if "mobile_results" in source else 4.6 if "web_production" in source else 3.35
    width = min(max_width, max_height / ratio)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    cap = document.add_paragraph(style="Caption")
    add_inline(cap, caption)


def parse_body(document: Document, markdown: str) -> tuple[int, int]:
    body = markdown.split("\n---\n", 1)[1]
    lines = body.splitlines()
    index = 0
    tables = figures = 0
    paragraph_lines: list[str] = []
    in_code = False
    code_language = ""
    code_lines: list[str] = []

    def flush():
        nonlocal paragraph_lines
        text = " ".join(line.strip() for line in paragraph_lines).strip()
        if text:
            figure_analysis = text.startswith("*Hình") and text.endswith("*")
            if figure_analysis:
                # The image alt text is already the numbered caption. Keep the
                # following italic sentence as analysis instead of duplicating it.
                analysis = re.sub(r"^Hình\s+\d+\.\s*", "Phân tích: ", text[1:-1])
                paragraph = document.add_paragraph()
                add_inline(paragraph, analysis)
                for run in paragraph.runs:
                    run.italic = True
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, text)
        paragraph_lines = []

    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("```"):
            flush()
            if in_code:
                style = "Output Block" if code_language in {"text", "output"} else "Code Block"
                paragraph = document.add_paragraph("\n".join(code_lines), style=style)
                paragraph.paragraph_format.keep_together = True
                code_lines = []
                code_language = ""
                in_code = False
            else:
                code_language = stripped[3:].strip().lower()
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(lines[index])
            index += 1
            continue
        if not stripped:
            flush()
            index += 1
            continue
        image = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image:
            flush()
            add_figure(document, image.group(2), image.group(1))
            figures += 1
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush()
            title = heading.group(2)
            if len(heading.group(1)) == 1 and (
                title.startswith("Phụ lục") or title == "Tài liệu tham khảo"
            ):
                document.add_page_break()
            document.add_heading(title, level=len(heading.group(1)))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1].strip()):
            flush()
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            if tables >= len(TABLE_CAPTIONS):
                raise RuntimeError("Missing table caption definition")
            add_table(document, rows, TABLE_CAPTIONS[tables])
            tables += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush()
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush()
    return tables, figures


def build() -> tuple[int, int]:
    generate_assets()
    document = Document()
    configure(document)
    document.core_properties.title = "Báo cáo Assignment 02 – From Data Representation to a Deployable Intelligent System"
    document.core_properties.author = "Nguyễn Thành Trung"
    document.core_properties.subject = "Diabetes, Vietnam House Price và E-commerce Customer Preference"
    add_cover(document)
    document.add_page_break()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    run = heading.add_run("MỤC LỤC")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Cập nhật mục lục trong Microsoft Word")
    document.add_page_break()
    tables, figures = parse_body(document, MARKDOWN.read_text(encoding="utf-8"))
    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.save(DOCX)
    return tables, figures


def validate(expected_tables: int, expected_figures: int) -> None:
    if not zipfile.is_zipfile(DOCX):
        raise RuntimeError("Invalid DOCX package")
    document = Document(DOCX)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    full = text + "\n" + table_text
    required = [
        "Nguyễn Thành Trung", "B23DCCN861", "D23CTPM01-B", "Dinh Que Tran",
        "768", "30.229", "568.454", "524.983", "120.000", "12.005",
        "0,7468", "0,6061", "1,2529", "1,6018", "0,4738", "27,36%",
        "0,9566", "0,9639", "0,9855", "0,9746", "0,9842",
        "không phải 47,38% Accuracy", "Linear SVC có F1 cao nhất",
        "Neo4j unavailable", "Vercel", "Render", "Expo Go",
    ]
    missing = [item for item in required if item not in full]
    if missing:
        raise RuntimeError(f"Missing report facts: {missing}")
    prose_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
        if paragraph.style.name not in {"Code Block", "Output Block"}
    )
    artifact_paragraphs = [
        paragraph.text for paragraph in document.paragraphs
        if paragraph.style.name not in {"Code Block", "Output Block"}
        and ("**" in paragraph.text or re.search(r"(?<!`)`[^`]+`", paragraph.text))
    ]
    if artifact_paragraphs:
        raise RuntimeError(f"Literal Markdown artifact in DOCX prose: {artifact_paragraphs[:3]}")
    images = [relationship for relationship in document.part.rels.values() if "image" in relationship.reltype]
    if len(document.tables) != expected_tables or len(images) != expected_figures:
        raise RuntimeError(f"Structure mismatch: tables={len(document.tables)}, images={len(images)}")
    duplicate_captions = [
        number for number in range(1, len(images) + 1)
        if text.count(f"Hình {number}.") != 1
    ]
    if duplicate_captions:
        raise RuntimeError(f"Missing or duplicated figure captions: {duplicate_captions}")
    bad_table_captions = [
        number for number in range(1, len(document.tables) + 1)
        if text.count(f"Bảng {number}.") != 1
    ]
    if bad_table_captions:
        raise RuntimeError(f"Missing or duplicated table captions: {bad_table_captions}")
    if "TOC \\o" not in document._element.xml or "PAGE" not in document.sections[0].footer._element.xml:
        raise RuntimeError("TOC/page-number field missing")
    external_links = {
        relationship.target_ref for relationship in document.part.rels.values()
        if relationship.reltype == RT.HYPERLINK and relationship.is_external
    }
    required_links = {
        "https://intelligent-system-assignment-01.vercel.app",
        "https://intelligent-system-assignment-01.onrender.com",
        "https://intelligent-system-assignment-01.onrender.com/docs",
        "https://github.com/Sagitoaz/intelligent-system-assignment-01",
    }
    if not required_links.issubset(external_links):
        raise RuntimeError(f"Missing clickable deployment links: {required_links - external_links}")
    section = document.sections[0]
    if abs(section.page_width.cm - 21) > 0.05 or abs(section.page_height.cm - 29.7) > 0.05:
        raise RuntimeError("Not A4")
    word_count = len(re.findall(r"\b\w+[\w-]*\b", text, flags=re.UNICODE))
    prose_word_count = len(re.findall(r"\b\w+[\w-]*\b", prose_text, flags=re.UNICODE))
    evidence_blocks = [
        paragraph for paragraph in document.paragraphs
        if paragraph.style.name in {"Code Block", "Output Block"}
    ]
    code_blocks = sum(paragraph.style.name == "Code Block" for paragraph in evidence_blocks)
    output_blocks = sum(paragraph.style.name == "Output Block" for paragraph in evidence_blocks)
    analysis_paragraphs = sum(
        "Phân tích" in paragraph.text for paragraph in document.paragraphs
    )
    evidence_lines = sum(paragraph.text.count("\n") + 1 for paragraph in evidence_blocks)
    max_evidence_line = max(
        (len(line) for paragraph in evidence_blocks for line in paragraph.text.splitlines()),
        default=0,
    )
    if max_evidence_line > 115:
        raise RuntimeError(f"Code/output line may overflow: {max_evidence_line} characters")
    content_width = section.page_width - section.left_margin - section.right_margin
    if any(shape.width > content_width for shape in document.inline_shapes):
        raise RuntimeError("Inline image exceeds the writable page width")
    page_breaks = document._element.xml.count('w:type="page"')
    estimated_pages = round(
        2
        + prose_word_count / 500
        + evidence_lines / 55
        + len(images) * 0.30
        + len(document.tables) * 0.16
    )
    print(f"DOCX={DOCX}")
    print(
        f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
        f"images={len(images)} words~={word_count} evidence_blocks={len(evidence_blocks)} "
        f"code_blocks={code_blocks} output_blocks={output_blocks} "
        f"analysis_paragraphs={analysis_paragraphs} evidence_lines={evidence_lines} "
        f"max_evidence_line={max_evidence_line} page_breaks={page_breaks}"
    )
    print(f"estimated_pages~={estimated_pages} valid_zip=True a4=True toc=True page_numbers=True facts=True")


if __name__ == "__main__":
    table_count, figure_count = build()
    validate(table_count, figure_count)
