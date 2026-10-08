"""Dựng báo cáo Word của tiểu luận từ các chương Markdown và kết quả thực nghiệm đã lưu.

Chạy:  .venv\\Scripts\\python.exe -m scripts.tieuluan.build_report
Sau đó: .venv\\Scripts\\python.exe -m scripts.tieuluan.finalize_docx   (Word cập nhật mục lục, số trang)

Quy ước trong file Markdown (docs/tieuluan/chapters/*.md):
- ``# `` chương (sang trang mới), ``## `` mục, ``### `` tiểu mục.
- ``![Hình 2.1. Chú thích.](../../figures/tieuluan/x.png){w=14}``: hình, chú thích đặt dưới hình.
- Dòng ``Bảng 2.1. Chú thích.`` đứng ngay trên bảng Markdown: chú thích bảng đặt trên bảng.
- ``$$ công thức LaTeX $$ (2.1)``: công thức hiển thị có đánh số.
- ``> `` khung "Hiểu nhanh"; ``[@khoa]`` trích dẫn; ``{{v:khoa}}`` số liệu; ``{{t:khoa}}`` bảng sinh tự động.
Mọi số liệu thực nghiệm trong văn bản đi qua ``{{v:...}}``/``{{t:...}}`` để luôn khớp với file kết quả.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from sklearn.metrics import roc_curve

from src.tieuluan.config import ANALYSIS_START
from src.tieuluan.experiment_data import OUT, RAW, ROOT

DOCS = ROOT / "docs/tieuluan"
CHAPTERS = DOCS / "chapters"
RESULTS = ROOT / "results/tieuluan/full"
FIGURES = ROOT / "figures/tieuluan"
EQUATIONS = FIGURES / "equations"
NAME = "tieuluan01_NhomLop_NhomTL_TrungNT"
SEEDS = (11, 22, 33)
MARKETS = ("sp500", "vnindex", "btc")
TABULAR = ("taiwan_bankruptcy", "credit_default")
LABEL = {"taiwan_bankruptcy": "Phá sản doanh nghiệp", "credit_default": "Vỡ nợ thẻ tín dụng",
         "sp500": "S&P 500", "vnindex": "VN-Index", "btc": "Bitcoin"}
KIND = {"mlp": "MLP", "logistic": "Hồi quy logistic", "random_forest": "Rừng ngẫu nhiên",
        "majority": "Cơ sở: luôn đoán lớp đa số", "persistence": "Cơ sở: lặp lại hướng phiên trước",
        "cnn4": "CNN4", "cnn8": "CNN8", "cnndeep": "CNN sâu (2 khối)", "rnn": "SimpleRNN", "lstm": "LSTM", "gru": "GRU"}
FW = {"scratch": "NumPy tự viết", "pytorch": "PyTorch", "keras": "Keras", "sklearn": "scikit-learn", "baseline": "—"}

# Bảng màu tham chiếu (3 ô đầu kiểm định an toàn cho người mù màu khi đặt cạnh nhau).
BLUE, ORANGE, AQUA, YELLOW = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
INK, INK2, MUTED, GRID = "#0b0b0b", "#52514e", "#8a8984", "#e4e3de"
FW_STYLE = {"scratch": (BLUE, "-", "o"), "pytorch": (ORANGE, "--", "s"), "keras": (AQUA, ":", "^")}
MARKET_COLOR = {"sp500": BLUE, "vnindex": ORANGE, "btc": AQUA}


# =============================================================================================
# Định dạng số kiểu Việt Nam
# =============================================================================================
def vn(x: float, nd: int = 3) -> str:
    s = f"{x:,.{nd}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".").replace("-", "−")


def sci(x: float) -> str:
    """Số rất nhỏ dạng a × 10⁻ᵏ (kiểu Việt Nam)."""
    if x == 0:
        return "0"
    exponent = int(np.floor(np.log10(abs(x))))
    mantissa = x / 10 ** exponent
    sup = str(exponent).translate(str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹"))
    return f"{vn(mantissa, 1)} × 10{sup}"


def vn_int(n) -> str:
    return f"{int(round(n)):,}".replace(",", ".")


def vn_pct(x: float, nd: int = 1) -> str:
    return vn(100 * x, nd) + "%"


def vn_date(s) -> str:
    d = pd.Timestamp(str(s))
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


# =============================================================================================
# Đọc kết quả
# =============================================================================================
class Results:
    def __init__(self):
        self.records = [json.loads(p.read_text(encoding="utf8")) for p in sorted(RESULTS.glob("*__*.json"))]
        if len(self.records) != 128:
            raise RuntimeError(f"Cần 128 bản ghi thực nghiệm, tìm thấy {len(self.records)}")
        self.analysis = json.loads((RESULTS / "analysis.json").read_text(encoding="utf8"))
        self.inventory = json.loads((OUT / "inventory.json").read_text(encoding="utf8"))
        self.quality = json.loads((ROOT / "data/tieuluan/data_quality.json").read_text(encoding="utf8"))
        self.environment = json.loads((RESULTS / "environment.json").read_text(encoding="utf8"))

    def runs(self, dataset, kind, framework=None):
        return [r for r in self.records if r["dataset"] == dataset and r["kind"] == kind
                and (framework is None or r["framework"] == framework)]

    def metric(self, dataset, kind, framework, name="roc_auc"):
        return np.array([r["metrics"][name] for r in self.runs(dataset, kind, framework)])

    def pred(self, dataset, kind, framework, seed):
        with np.load(RESULTS / f"{dataset}__{kind}__{framework}__{seed}_predictions.npz", allow_pickle=False) as d:
            return {k: d[k] for k in d.files}

    def ensemble(self, dataset, kind, framework):
        runs = [self.pred(dataset, kind, framework, s) for s in SEEDS]
        return runs[0]["y"], np.mean([r["probability"] for r in runs], axis=0)


def mean_sd(values, nd=3):
    values = np.asarray(values, dtype=float)
    return f"{vn(values.mean(), nd)} ± {vn(values.std(ddof=1), nd)}" if len(values) > 1 else vn(values[0], nd)


def ci_text(ci, nd=3):
    return f"{vn(ci['auc'], nd)} [{vn(ci['low'], nd)}; {vn(ci['high'], nd)}]"


# =============================================================================================
# Số liệu {{v:...}} và bảng {{t:...}} sinh từ kết quả
# =============================================================================================
def market_frame(key):
    f = pd.read_csv(RAW / "indices" / f"{key}.csv", parse_dates=["date"])
    return f[f.date >= pd.Timestamp(ANALYSIS_START[key])].reset_index(drop=True)


def compute_values(res: Results) -> dict:
    v = {}
    app_assets = json.loads((ROOT / 'data/tieuluan/app/metadata.json').read_text(encoding='utf-8'))
    v['app.credit.threshold'] = vn_pct(app_assets['datasets']['credit_default']['threshold'])
    app_check = json.loads((ROOT / 'data/tieuluan/app/verification.json').read_text(encoding='utf-8'))
    v['app.peak_memory_mb'] = vn(app_check['peak_working_set_mb'], 1)
    inv, an = res.inventory, res.analysis
    for key in MARKETS + TABULAR:
        s = inv[key]["splits"]
        for split in ("train", "val", "test"):
            v[f"n.{key}.{split}"] = vn_int(s[split]["n"])
            v[f"pos.{key}.{split}"] = vn_pct(s[split]["positive"] / s[split]["n"])
            v[f"npos.{key}.{split}"] = vn_int(s[split]["positive"])
        v[f"n.{key}.total"] = vn_int(sum(s[x]["n"] for x in ("train", "val", "test")))
    vol = {}
    for key in MARKETS:
        f = market_frame(key)
        v[f"rows.{key}"] = vn_int(len(f))
        v[f"first.{key}"] = vn_date(f.date.iloc[0])
        v[f"last.{key}"] = vn_date(f.date.iloc[-1])
        v[f"raw_rows.{key}"] = vn_int(inv[key]["raw_rows"])
        r = np.log(f.close).diff().dropna()
        vol[key] = r.std() * np.sqrt(365 if key == "btc" else 252)
        v[f"vol.{key}"] = vn_pct(vol[key])
        # Tự tương quan bậc 1 của lợi suất ngày: > 0 nghĩa là phiên tăng hay được nối tiếp bởi phiên tăng
        v[f"ac1.{key}"] = vn(float(r.autocorr(1)), 3)
        up = (r > 0).to_numpy()
        v[f"pupup.{key}"] = vn_pct(up[1:][up[:-1]].mean())
        v[f"pupdown.{key}"] = vn_pct(up[1:][~up[:-1]].mean())
        # Mẫu đầu ra của Chương 4: cửa sổ 20 lợi suất đầu tiên của tập test (khôi phục về đơn vị lợi suất gốc)
        with np.load(OUT / f"{key}.npz", allow_pickle=False) as d:
            first = d["test_sequence"][0, :, 0] * d["scaler_scale"][0] + d["scaler_mean"][0]
            v[f"seqhead.{key}"] = "; ".join(vn(100 * float(x), 2) + "%" for x in first[:5])
            v[f"seqdate.{key}"] = vn_date(d["test_date"][0])
            v[f"seqnext.{key}"] = vn(100 * float(d["test_return"][0]), 2) + "%"
            v[f"seqlabel.{key}"] = "tăng" if float(d["test_y"][0]) == 1 else "không tăng"
    v["vol_ratio"] = vn(vol["btc"] / max(vol["sp500"], vol["vnindex"]), 1)
    for key in TABULAR:
        v[f"rows.{key}"] = vn_int(inv[key]["raw_rows"])
        v[f"features.{key}"] = str(inv[key]["features"])
        v[f"positive.{key}"] = vn_int(inv[key]["positive"])
        v[f"positive_rate.{key}"] = vn_pct(inv[key]["positive"] / inv[key]["raw_rows"], 2)
        v[f"majority_acc.{key}"] = vn_pct(1 - inv[key]["positive"] / inv[key]["raw_rows"], 2)
    q = res.quality["vnindex_cross_source"]
    for k in ("rows_ssi", "rows_dnse", "rows_vndirect", "overlap_ssi_dnse", "close_disagreements_ssi_vs_dnse",
              "vndirect_votes_for_ssi", "vndirect_votes_for_dnse", "dnse_rows_filling_ssi_gaps_after_cut",
              "overlap_ssi_vndirect", "close_disagreements_ssi_vs_vndirect", "rows_final"):
        v[f"vnq.{k}"] = vn_int(q[k])
    v["vnq.agree_pct"] = vn_pct(1 - q["close_disagreements_ssi_vs_dnse"] / q["overlap_ssi_dnse"], 2)
    # Số phiên SSI bỏ trống trong đợt thiếu dữ liệu 07–08/2009 (được bù từ DNSE)
    v["vnq.gap2009"] = str(sum("2009-07-17" < d < "2009-08-25" for d in q["dnse_fill_dates_after_cut"]))
    # Kết quả theo (bộ dữ liệu, mô hình, cài đặt)
    for r in res.records:
        key = f"{r['dataset']}.{r['kind']}.{r['framework']}"
        if f"auc.{key}" in v:
            continue
        for name, short in (("roc_auc", "auc"), ("average_precision", "ap"), ("balanced_accuracy", "bacc"), ("f1", "f1"),
                            ("recall", "recall"), ("precision", "precision"), ("accuracy", "acc")):
            vals = res.metric(r["dataset"], r["kind"], r["framework"], name)
            v[f"{short}.{key}"] = vn(vals.mean())
            v[f"{short}sd.{key}"] = mean_sd(vals)
            v[f"{short}pct.{key}"] = vn_pct(vals.mean())
        times = [x["training_seconds"] for x in res.runs(r["dataset"], r["kind"], r["framework"])]
        v[f"time.{key}"] = vn(float(np.mean(times)), 1)
        v[f"params.{key}"] = vn_int(r["parameters"]) if r.get("parameters") else "—"
    for key in MARKETS:
        for kind, ci in an["auc_ci"][key].items():
            v[f"ci.{key}.{kind}"] = ci_text(ci)
            v[f"cilow.{key}.{kind}"] = vn(ci["low"])
            v[f"cihigh.{key}.{kind}"] = vn(ci["high"])
            v[f"ciauc.{key}.{kind}"] = vn(ci["auc"])
        for strat, b in an["backtest"][key].items():
            v[f"bt.{key}.{strat}.ret"] = vn_pct(b["annual_return"])
            v[f"bt.{key}.{strat}.sharpe"] = vn(b["sharpe"], 2)
            v[f"bt.{key}.{strat}.mdd"] = vn_pct(b["max_drawdown"])
            v[f"bt.{key}.{strat}.exposure"] = vn_pct(b["exposure"], 0)
            v[f"bt.{key}.{strat}.entries"] = vn_int(b["entries"])
            v[f"bt.{key}.{strat}.cost"] = vn_pct(b["total_cost"])
            v[f"bt.{key}.{strat}.retgross"] = vn_pct(b["annual_return_gross"])
            v[f"bt.{key}.{strat}.sharpegross"] = vn(b["sharpe_gross"], 2)
    for key in TABULAR:
        for name, ci in an["uci"][key].items():
            if "auc" in ci:
                v[f"uci.{key}.{name}"] = ci_text(ci)
            else:
                v[f"ucidiff.{key}.{name}"] = f"{vn(ci['difference'])} [{vn(ci['low'])}; {vn(ci['high'])}]"
    for key in ("sp500", "vnindex", "btc"):
        v[f"cost.{key}"] = vn_pct(an["assumptions"]["cost_per_side"][key], 2)
    ver = json.loads((ROOT / "results/tieuluan/verification.json").read_text(encoding="utf8"))
    for name, item in ver["gradient_check"].items():
        v[f"gc.{name}"] = sci(item["max_relative_error"])
        if "norm_relative_error" in item:
            v[f"gcnorm.{name}"] = sci(item["norm_relative_error"])
    for kind, diffs in ver["forward_parity_max_abs_logit_diff"].items():
        v[f"parity.{kind}"] = sci(max(diffs.values()))
    for key, kind in [(k, "mlp") for k in TABULAR] + [(k, kd) for k in MARKETS for kd in ("cnn4", "lstm")]:
        base = res.pred(key, kind, "scratch", 11)["probability"]
        for fw in ("pytorch", "keras"):
            v[f"parity_trained.{key}.{kind}.{fw}"] = sci(float(np.max(np.abs(res.pred(key, kind, fw, 11)["probability"] - base))))
    # Chênh lệch ROC-AUC lớn nhất (theo từng seed) giữa bản tự viết và từng thư viện
    for kind, datasets in (("mlp", TABULAR), ("cnn4", MARKETS), ("lstm", MARKETS)):
        for fw in ("pytorch", "keras"):
            diff = max(abs(r["metrics"]["roc_auc"] - s["metrics"]["roc_auc"])
                       for key in datasets for r in res.runs(key, kind, fw) for s in res.runs(key, kind, "scratch")
                       if r["seed"] == s["seed"])
            v[f"aucdiff.{kind}.{fw}"] = f"tối đa {vn(diff, 4)}" if diff >= 5e-5 else "dưới 0,0001"
    for key in MARKETS:
        with np.load(OUT / f"{key}.npz", allow_pickle=False) as d:
            dates = pd.to_datetime(d["test_target_date"])
        f = market_frame(key).set_index("date")["close"]
        actual = f.reindex(dates).to_numpy()
        previous = f.shift(1).reindex(dates).to_numpy()
        r2 = 1 - np.sum((actual - previous) ** 2) / np.sum((actual - actual.mean()) ** 2)
        v[f"naive_r2.{key}"] = vn(r2, 4)
    sens_path = ROOT / "results/tieuluan/sensitivity_tabular.json"
    if sens_path.exists():
        sens = json.loads(sens_path.read_text(encoding="utf8"))
        for key, models in sens.items():
            for name, item in models.items():
                v[f"sens.{key}.{name}"] = vn(item["auc_mean"])
                v[f"sensci.{key}.{name}"] = ci_text(item["ensemble_ci"])
    audit = json.loads((ROOT / "results/tieuluan/audit.json").read_text(encoding="utf8"))
    v["audit.records"] = str(audit["records_recomputed"])
    v["audit.reloaded"] = str(audit["checkpoints_reloaded"])
    err = audit["max_probability_reload_error"]
    v["audit.match"] = "trùng khớp tuyệt đối với" if err == 0 else f"lệch tối đa {sci(err)} so với"
    ablation_path = ROOT / "results/tieuluan/ablation_gasf.json"
    if ablation_path.exists():
        for key, item in json.loads(ablation_path.read_text(encoding="utf8")).items():
            v[f"abl.{key}"] = vn(item["auc_mean"])
            v[f"ablci.{key}"] = ci_text(item["ensemble_ci"])
    env = res.environment
    for k in ("python", "numpy", "torch", "tensorflow", "keras", "sklearn"):
        v[f"env.{k}"] = env[k].split("+")[0]
    v["env.epochs"], v["env.patience"], v["env.batch"] = str(env["epochs"]), str(env["patience"]), str(env["batch_size"])
    v["n_records"] = str(len(res.records))
    v["n_trained"] = str(sum(r["framework"] != "baseline" for r in res.records))
    return v


def md_table(caption, headers, rows):
    lines = [caption, "| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    lines += ["| " + " | ".join(str(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def compute_tables(res: Results, v: dict) -> dict:
    t = {}
    inv = res.inventory
    # ---- Chương 1
    rows = []
    for key in MARKETS:
        f = market_frame(key)
        r = np.log(f.close).diff().dropna()
        ppy = 365 if key == "btc" else 252
        rows.append([LABEL[key], f"{vn_date(f.date.iloc[0])} – {vn_date(f.date.iloc[-1])}", vn_int(len(f)),
                     vn(100 * r.mean(), 3), vn(100 * r.std(), 2), vn(100 * r.std() * np.sqrt(ppy), 1),
                     vn(100 * r.min(), 1), vn(100 * r.max(), 1), vn(float(r.kurt()), 1), vn_pct((r > 0).mean())])
    t["ch1_market_stats"] = md_table(
        "Bảng {n}. Thống kê mô tả lợi suất logarit theo ngày trong giai đoạn phân tích (tính từ dữ liệu của tiểu luận).",
        ["Chỉ số", "Giai đoạn", "Số phiên", "TB (%)", "Độ lệch chuẩn (%)", "Biến động năm (%)", "Nhỏ nhất (%)",
         "Lớn nhất (%)", "Độ nhọn dư", "Phiên tăng"], rows)
    rows = []
    for key in MARKETS:
        f = market_frame(key)
        for _, x in f.head(2).iterrows():
            rows.append([LABEL[key], vn_date(x.date), vn(x.open, 2), vn(x.high, 2), vn(x.low, 2), vn(x.close, 2),
                         vn_int(x.volume)])
    t["ch1_market_samples"] = md_table("Bảng {n}. Hai dòng dữ liệu gốc đầu tiên của mỗi chuỗi trong giai đoạn phân tích.",
                                       ["Chỉ số", "Ngày", "Mở cửa", "Cao nhất", "Thấp nhất", "Đóng cửa", "Khối lượng"], rows)
    # ---- Chương 2
    t["ch2_datasets"] = md_table(
        "Bảng {n}. Hai bộ dữ liệu bảng của Chương 2 (số liệu kiểm kê trực tiếp từ file tải về).",
        ["Thuộc tính", LABEL["taiwan_bankruptcy"], LABEL["credit_default"]],
        [["Nguồn", "UCI, mã 572 [@uci572]", "UCI, mã 350 [@uci350]"],
         ["Đơn vị quan sát", "Một doanh nghiệp niêm yết (Đài Loan, 1999–2009)", "Một khách hàng thẻ tín dụng (Đài Loan, 2005)"],
         ["Số dòng × số đặc trưng", f"{v['rows.taiwan_bankruptcy']} × {v['features.taiwan_bankruptcy']}",
          f"{v['rows.credit_default']} × {v['features.credit_default']}"],
         ["Nhãn dương (lớp 1)", f"Phá sản: {v['positive.taiwan_bankruptcy']} ({v['positive_rate.taiwan_bankruptcy']})",
          f"Vỡ nợ tháng sau: {v['positive.credit_default']} ({v['positive_rate.credit_default']})"],
         ["Ô dữ liệu thiếu", "0", "0"], ["Giấy phép", "CC BY 4.0", "CC BY 4.0"]])
    tw = pd.read_csv(RAW / "uci/taiwan_bankruptcy/data.csv")
    cd = pd.read_excel(next((RAW / "uci/credit_default").glob("*.xls")), header=1)
    col = {name: next(c for c in tw.columns if name in c) for name in ("ROA(C)", "Debt ratio", "Net Income to Total Assets")}
    rows = []
    for i in range(3):
        rows.append([LABEL["taiwan_bankruptcy"], str(i + 1), str(int(tw.iloc[i, 0])),
                     f"ROA(C) = {vn(float(tw.loc[i, col['ROA(C)']]), 4)}; tỷ lệ nợ = {vn(float(tw.loc[i, col['Debt ratio']]), 4)}; "
                     f"lãi ròng/tổng tài sản = {vn(float(tw.loc[i, col['Net Income to Total Assets']]), 4)}"])
    for i in range(3):
        x = cd.iloc[i]
        rows.append([LABEL["credit_default"], str(i + 1), str(int(x[cd.columns[-1]])),
                     f"hạn mức = {vn_int(x['LIMIT_BAL'])} TWD; tuổi = {int(x['AGE'])}; PAY_0 = {int(x['PAY_0'])}; "
                     f"dư nợ tháng 9 = {vn_int(x['BILL_AMT1'])} TWD"])
    t["ch2_samples"] = md_table("Bảng {n}. Ba dòng đầu của mỗi bộ dữ liệu (trích ba – bốn đặc trưng; nhãn 1 là có rủi ro).",
                                ["Bộ dữ liệu", "Dòng", "Nhãn", "Một số đặc trưng"], rows)
    sens_path = ROOT / "results/tieuluan/sensitivity_tabular.json"
    if sens_path.exists():
        sens = json.loads(sens_path.read_text(encoding="utf8"))
        rows = []
        for key in TABULAR:
            for name, kind, fw in (("logistic", "logistic", "sklearn"), ("random_forest", "random_forest", "sklearn"), ("mlp", "mlp", "pytorch")):
                rows.append([LABEL[key], KIND[kind], v[f"auc.{key}.{kind}.{fw}"], vn(sens[key][name]["auc_mean"]),
                             ci_text(sens[key][name]["ensemble_ci"])])
        t["ch2_sensitivity"] = md_table("Bảng {n}. Độ nhạy theo tiền xử lý: ROC-AUC trên test khi chuẩn hóa trực tiếp (thực nghiệm chính) và khi biến đổi log có dấu trước khi chuẩn hóa.",
                                        ["Bộ dữ liệu", "Mô hình", "AUC – chuẩn hóa trực tiếp", "AUC – log rồi chuẩn hóa", "AUC log [CI 95%]"], rows)
    rows = []
    for key in TABULAR:
        s = inv[key]["splits"]
        rows.append([LABEL[key]] + [f"{v[f'n.{key}.{x}']} ({v[f'pos.{key}.{x}']})" for x in ("train", "val", "test")])
    t["ch2_splits"] = md_table("Bảng {n}. Kích thước các tập sau khi chia phân tầng 60/20/20 (trong ngoặc: tỉ lệ lớp dương).",
                               ["Bộ dữ liệu", "Train", "Validation", "Test"], rows)
    t["ch2_frameworks"] = framework_table(res, "Bảng {n}. Cùng một MLP cài đặt bằng ba cách (trung bình ± độ lệch chuẩn qua 3 seed, đánh giá trên test).",
                                          TABULAR, "mlp")
    rows = []
    for key in TABULAR:
        for kind, fw in (("logistic", "sklearn"), ("random_forest", "sklearn"), ("mlp", "pytorch"), ("majority", "baseline")):
            ci = res.analysis["uci"][key].get(kind) if kind != "majority" else None
            rows.append([LABEL[key], KIND[kind], v[f"aucsd.{key}.{kind}.{fw}"], ci_text(ci) if ci else "—",
                         v[f"apsd.{key}.{kind}.{fw}"], v[f"baccsd.{key}.{kind}.{fw}"], v[f"recallsd.{key}.{kind}.{fw}"]])
    t["ch2_algorithms"] = md_table(
        "Bảng {n}. So sánh ba thuật toán trên tập test. CI 95%: bootstrap phân tầng 2.000 lần cho dự báo trung bình 3 seed.",
        ["Bộ dữ liệu", "Mô hình", "ROC-AUC", "AUC [CI 95%]", "AP", "Balanced acc.", "Recall lớp 1"], rows)
    # ---- Chương 3 & 4
    rows = []
    for key in MARKETS:
        rows.append([LABEL[key]] + [f"{v[f'n.{key}.{x}']} ({v[f'pos.{key}.{x}']})" for x in ("train", "val", "test")])
    split_caption = "Bảng {n}. Số mẫu (cửa sổ 20 phiên) của mỗi tập; trong ngoặc là tỉ lệ mẫu có nhãn “tăng”."
    t["ch3_splits"] = md_table(split_caption, ["Chỉ số", "Train (≤ 2019)", "Validation (2020–2022)", "Test (2023–09/2026)"], rows)
    t["ch3_frameworks"] = framework_table(res, "Bảng {n}. Cùng một CNN4 cài đặt bằng ba cách (trung bình ± độ lệch chuẩn qua 3 seed, tập test).",
                                          MARKETS, "cnn4")
    t["ch3_architectures"] = architecture_table(res, "Bảng {n}. So sánh kiến trúc CNN với hai đường cơ sở (PyTorch, 3 seed). CI 95%: bootstrap theo khối 20 phiên cho dự báo trung bình 3 seed.",
                                                ("cnn4", "cnn8", "cnndeep"))
    t["ch4_frameworks"] = framework_table(res, "Bảng {n}. Cùng một LSTM cài đặt bằng ba cách (trung bình ± độ lệch chuẩn qua 3 seed, tập test).",
                                          MARKETS, "lstm")
    t["ch4_architectures"] = architecture_table(res, "Bảng {n}. So sánh SimpleRNN, LSTM và GRU với hai đường cơ sở (PyTorch, 3 seed). CI 95%: bootstrap theo khối 20 phiên.",
                                                ("rnn", "lstm", "gru"))
    rows = []
    names = {"buy_hold": "Mua và giữ", "persistence": "Lặp lại hướng phiên trước", "cnn4": "CNN4", "cnn8": "CNN8",
             "lstm": "LSTM", "gru": "GRU"}
    for key in MARKETS:
        for strat in ("buy_hold", "persistence", "cnn4", "cnn8", "lstm", "gru"):
            b = res.analysis["backtest"][key][strat]
            rows.append([LABEL[key], names[strat], vn_pct(b["annual_return_gross"]), vn_pct(b["annual_return"]),
                         vn(b["sharpe"], 2), vn_pct(b["annual_volatility"]), vn_pct(b["max_drawdown"]),
                         vn_pct(b["exposure"], 0), vn_int(b["entries"])])
    t["ch4_backtest"] = md_table(
        "Bảng {n}. Backtest minh họa trên tập test (01/2023–09/2026); tín hiệu từ dự báo trung bình 3 seed, ngưỡng chọn trên validation. Các cột sau “Lợi suất năm trước phí” đều đã trừ phí giao dịch.",
        ["Chỉ số", "Chiến lược", "Lợi suất năm trước phí", "Lợi suất năm sau phí", "Sharpe", "Biến động năm",
         "Sụt giảm tối đa", "Thời gian nắm giữ", "Số lần mua"], rows)
    return t


def framework_table(res, caption, datasets, kind):
    rows = []
    for key in datasets:
        for fw in ("scratch", "pytorch", "keras"):
            runs = res.runs(key, kind, fw)
            auc = [r["metrics"]["roc_auc"] for r in runs]
            ap = [r["metrics"]["average_precision"] for r in runs]
            bacc = [r["metrics"]["balanced_accuracy"] for r in runs]
            rows.append([LABEL[key], FW[fw], mean_sd(auc), mean_sd(ap), mean_sd(bacc), vn_int(runs[0]["parameters"]),
                         "/".join(str(r["best_epoch"]) for r in sorted(runs, key=lambda r: r["seed"])),
                         vn(float(np.mean([r["training_seconds"] for r in runs])), 1)])
    return md_table(caption, ["Dữ liệu", "Cài đặt", "ROC-AUC", "AP", "Balanced acc.", "Tham số", "Epoch tốt nhất", "Giây/lượt"], rows)


def architecture_table(res, caption, kinds):
    rows = []
    for key in MARKETS:
        for kind in kinds + ("persistence", "majority"):
            fw = "baseline" if kind in ("persistence", "majority") else "pytorch"
            runs = res.runs(key, kind, fw)
            auc = [r["metrics"]["roc_auc"] for r in runs]
            bacc = [r["metrics"]["balanced_accuracy"] for r in runs]
            ci = res.analysis["auc_ci"][key].get(kind)
            rows.append([LABEL[key], KIND[kind], vn_int(runs[0]["parameters"]) if runs[0]["parameters"] else "0",
                         mean_sd(auc), ci_text(ci) if ci else "—", mean_sd(bacc)])
    return md_table(caption, ["Chỉ số", "Mô hình", "Tham số", "ROC-AUC (3 seed)", "AUC [CI 95%]", "Balanced acc."], rows)


# =============================================================================================
# Hình từ dữ liệu và kết quả
# =============================================================================================
def style_axes(ax):
    ax.grid(color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#b9b8b2")
    ax.tick_params(colors=INK2, labelsize=7.5)


def savefig(fig, name):
    FIGURES.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIGURES / name, dpi=220, bbox_inches="tight", pad_inches=0.04, facecolor="white")
    plt.close(fig)


def make_figures(res: Results):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8, "text.color": INK, "axes.labelcolor": INK2,
                         "axes.titlesize": 8.5, "axes.titleweight": "bold", "legend.fontsize": 7.2, "legend.frameon": False})
    # 1) Lịch sử giá và ranh giới chia tập
    fig, axs = plt.subplots(3, 1, figsize=(6.8, 5.0), sharex=False)
    for ax, key in zip(axs, MARKETS):
        f = pd.read_csv(RAW / "indices" / f"{key}.csv", parse_dates=["date"])
        start = pd.Timestamp(ANALYSIS_START[key])
        used, excluded = f[f.date >= start], f[f.date < start]
        if len(excluded):
            ax.plot(excluded.date, excluded.close, color=MUTED, lw=0.8)
        ax.plot(used.date, used.close, color=MARKET_COLOR[key], lw=0.9)
        ax.set_yscale("log")
        for a, b, label in [(used.date.iloc[0], pd.Timestamp("2019-12-31"), "Train"),
                            (pd.Timestamp("2020-01-01"), pd.Timestamp("2022-12-31"), "Validation"),
                            (pd.Timestamp("2023-01-01"), pd.Timestamp("2026-09-30"), "Test")]:
            ax.axvspan(a, b, color={"Train": "#eef4fc", "Validation": "#fdf3e6", "Test": "#e9f7f1"}[label], zorder=0)
            ax.text(a + (b - a) / 2, 1.02, label, transform=ax.get_xaxis_transform(), ha="center", va="bottom",
                    fontsize=6.8, color=INK2)
        title = LABEL[key] + (" — đoạn xám 2000–2009 không dùng (giá nguồn làm tròn số nguyên)" if len(excluded) else "")
        ax.set_title(title, loc="left", fontsize=8, pad=11)
        ax.set_ylabel("Mức giá (thang log)", fontsize=7.2)
        style_axes(ax)
        ax.set_xlim(pd.Timestamp("2000-01-01"), pd.Timestamp("2026-12-31"))
    fig.tight_layout(h_pad=1.2)
    savefig(fig, "ch1_market_history.png")

    # 2) Phân phối lợi suất ngày so với phân phối chuẩn cùng trung bình/độ lệch
    fig, axs = plt.subplots(1, 3, figsize=(6.8, 2.2), sharey=True)
    for ax, key in zip(axs, MARKETS):
        r = 100 * np.log(market_frame(key).close).diff().dropna()
        lim = 12 if key == "btc" else 7
        bins = np.linspace(-lim, lim, 61)
        ax.hist(r.clip(-lim, lim), bins=bins, density=True, color=MARKET_COLOR[key], alpha=0.85, edgecolor="white", lw=0.3)
        xs = np.linspace(-lim, lim, 300)
        ax.plot(xs, np.exp(-(xs - r.mean()) ** 2 / (2 * r.var())) / np.sqrt(2 * np.pi * r.var()), color=INK, lw=1, ls="--")
        ax.set_title(LABEL[key], fontsize=8)
        ax.set_xlabel("Lợi suất ngày (%)", fontsize=7.2)
        style_axes(ax)
    axs[0].set_ylabel("Mật độ", fontsize=7.2)
    axs[-1].plot([], [], color=INK, lw=1, ls="--", label="Phân phối chuẩn\ncùng TB, độ lệch")
    axs[-1].legend(loc="upper right")
    fig.tight_layout()
    savefig(fig, "ch1_return_distribution.png")

    # 3) Hai bộ dữ liệu bảng: phân bố lớp và một đặc trưng tiêu biểu
    fig, axs = plt.subplots(1, 3, figsize=(6.8, 2.3), gridspec_kw={"width_ratios": [1.05, 1, 1]})
    ax = axs[0]
    labels, neg, pos = [], [], []
    for key in TABULAR:
        n, p = res.inventory[key]["raw_rows"], res.inventory[key]["positive"]
        labels.append(LABEL[key].replace(" ", "\n", 1))
        neg.append(n - p)
        pos.append(p)
    yy = np.arange(len(labels))
    ax.barh(yy, neg, color="#cdd9ea", height=0.55, label="Lớp 0")
    ax.barh(yy, pos, left=neg, color=ORANGE, height=0.55, label="Lớp 1 (rủi ro)")
    for i in range(len(labels)):
        ax.text(neg[i] + pos[i] + 300, i, f"{vn_pct(pos[i] / (neg[i] + pos[i]))}", va="center", fontsize=7, color=INK2)
    ax.set_yticks(yy, labels, fontsize=7)
    ax.set_xlabel("Số dòng", fontsize=7.2)
    ax.set_title("Phân bố lớp", fontsize=8)
    ax.legend(loc="lower right", fontsize=6.5)
    ax.set_xlim(0, 36000)
    style_axes(ax)
    tw = pd.read_csv(RAW / "uci/taiwan_bankruptcy/data.csv")
    col = [c for c in tw.columns if "Net Income to Total Assets" in c][0]
    for cls, color, lab in [(0, "#9fb8d9", "Không phá sản"), (1, ORANGE, "Phá sản")]:
        axs[1].hist(tw.loc[tw.iloc[:, 0] == cls, col], bins=np.linspace(0.6, 0.9, 40), density=True, color=color,
                    alpha=0.75, label=lab, edgecolor="white", lw=0.3)
    axs[1].set_title("Phá sản: lãi ròng / tổng tài sản", fontsize=8)
    axs[1].set_xlabel("Giá trị (đã chuẩn hóa trong nguồn)", fontsize=7)
    axs[1].legend(fontsize=6.5, loc="upper left")
    style_axes(axs[1])
    cd = pd.read_excel(next((RAW / "uci/credit_default").glob("*.xls")), header=1)
    target = cd.columns[-1]
    pay = cd["PAY_0"].clip(-2, 4)
    rate = cd.groupby(pay)[target].mean()
    count = cd.groupby(pay)[target].size()
    axs[2].bar(rate.index.astype(int), rate.values, color=[ORANGE if x >= 1 else "#9fb8d9" for x in rate.index], width=0.7)
    for x, y, n in zip(rate.index, rate.values, count.values):
        axs[2].text(x, y + 0.02, vn_int(n), ha="center", fontsize=5.8, color=INK2)
    axs[2].set_xticks(range(-2, 5), ["−2", "−1", "0", "1", "2", "3", "≥4"], fontsize=7)
    axs[2].set_title("Vỡ nợ theo trạng thái trả nợ PAY_0", fontsize=8)
    axs[2].set_xlabel("PAY_0 (≥ 1: số tháng trễ hạn)", fontsize=7)
    axs[2].set_ylabel("Tỉ lệ vỡ nợ", fontsize=7)
    axs[2].set_ylim(0, 1)
    style_axes(axs[2])
    fig.tight_layout()
    savefig(fig, "ch2_data_overview.png")

    # 4) Đường học validation của mô hình đại diện ở mỗi chương (seed 11)
    for chapter, kind, datasets in [(2, "mlp", TABULAR), (3, "cnn4", MARKETS), (4, "lstm", MARKETS)]:
        fig, axs = plt.subplots(1, len(datasets), figsize=(6.8, 2.1), squeeze=False)
        for ax, key in zip(axs[0], datasets):
            for fw in ("scratch", "pytorch", "keras"):
                r = next(x for x in res.runs(key, kind, fw) if x["seed"] == 11)
                color, ls, marker = FW_STYLE[fw]
                ep = [h["epoch"] for h in r["history"]]
                ax.plot(ep, [h["val_loss"] for h in r["history"]], color=color, ls=ls, lw=1.5, marker=marker, ms=2.6,
                        markevery=max(1, len(ep) // 8), label=FW[fw])
                ax.axvline(r["best_epoch"], color=MUTED, lw=0.6, ls=":")
            ax.set_title(LABEL[key], fontsize=8)
            ax.set_xlabel("Epoch", fontsize=7.2)
            style_axes(ax)
        axs[0, 0].set_ylabel("BCE trên validation", fontsize=7.2)
        axs[0, -1].legend(loc="upper right")
        fig.tight_layout()
        savefig(fig, f"ch{chapter}_learning_curves.png")

    # 5) ROC của ba thuật toán trên hai bộ dữ liệu bảng
    fig, axs = plt.subplots(1, 2, figsize=(6.0, 2.7))
    for ax, key in zip(axs, TABULAR):
        curves = [("Hồi quy logistic", res.pred(key, "logistic", "sklearn", 11), BLUE, "-"),
                  ("Rừng ngẫu nhiên", None, ORANGE, "--"), ("MLP (TB 3 seed)", None, AQUA, ":")]
        y_rf, p_rf = res.ensemble(key, "random_forest", "sklearn")
        y_mlp, p_mlp = res.ensemble(key, "mlp", "pytorch")
        for label, pred, color, ls in curves:
            y, p = (pred["y"], pred["probability"]) if pred is not None else ((y_rf, p_rf) if "Rừng" in label else (y_mlp, p_mlp))
            fpr, tpr, _ = roc_curve(y, p)
            ax.plot(fpr, tpr, color=color, ls=ls, lw=1.6, label=label)
        ax.plot([0, 1], [0, 1], color=MUTED, lw=0.8, ls="--")
        ax.set_title(LABEL[key], fontsize=8)
        ax.set_xlabel("Tỉ lệ dương giả (FPR)", fontsize=7.2)
        ax.set_aspect("equal")
        style_axes(ax)
    axs[0].set_ylabel("Tỉ lệ dương thật (TPR)", fontsize=7.2)
    axs[1].legend(loc="lower right")
    fig.tight_layout()
    savefig(fig, "ch2_roc.png")

    # 6) Ảnh GASF minh họa
    fig, axs = plt.subplots(2, 3, figsize=(6.8, 4.1), gridspec_kw={"height_ratios": [0.75, 1]})
    for col, key in enumerate(MARKETS):
        with np.load(OUT / f"{key}.npz", allow_pickle=False) as d:
            image, date = d["test_image"][0, 0], str(d["test_date"][0])
        f = market_frame(key)
        idx = int(np.flatnonzero(f.date.dt.strftime("%Y-%m-%d") == date)[0])
        window = f.close.iloc[idx - 19:idx + 1].to_numpy()
        axs[0, col].plot(range(1, 21), window, color=MARKET_COLOR[key], lw=1.4, marker="o", ms=2.2)
        axs[0, col].set_title(f"{LABEL[key]}\n20 phiên đến {vn_date(date)}", fontsize=7.6)
        axs[0, col].set_xticks([1, 5, 10, 15, 20])
        style_axes(axs[0, col])
        im = axs[1, col].imshow(image, cmap="RdBu_r", vmin=-1, vmax=1)
        axs[1, col].set_xticks([0, 9, 19], ["1", "10", "20"])
        axs[1, col].set_yticks([0, 9, 19], ["1", "10", "20"])
        axs[1, col].tick_params(labelsize=6.8, colors=INK2)
        axs[1, col].set_xlabel("Phiên j", fontsize=7)
    axs[1, 0].set_ylabel("Phiên i", fontsize=7)
    axs[0, 0].set_ylabel("Giá đóng cửa", fontsize=7)
    cbar = fig.colorbar(im, ax=axs[1, :], fraction=0.025, pad=0.02)
    cbar.ax.tick_params(labelsize=6.5)
    cbar.set_label("GASF = cos(φᵢ + φⱼ)", fontsize=7)
    savefig(fig, "ch3_gasf_examples.png")

    # 7) AUC kèm khoảng tin cậy 95% cho các kiến trúc ở mỗi thị trường
    for chapter, kinds in [(3, ("cnn4", "cnn8", "cnndeep", "persistence")), (4, ("rnn", "lstm", "gru", "persistence"))]:
        fig, axs = plt.subplots(1, 3, figsize=(6.8, 2.2), sharey=True)
        for ax, key in zip(axs, MARKETS):
            for i, kind in enumerate(kinds):
                ci = res.analysis["auc_ci"][key][kind]
                color = MUTED if kind == "persistence" else MARKET_COLOR[key]
                ax.plot([ci["low"], ci["high"]], [i, i], color=color, lw=2.2, solid_capstyle="round")
                ax.plot(ci["auc"], i, "o", color=color, ms=5, mec="white", mew=1)
            ax.axvline(0.5, color=INK2, lw=0.8, ls="--")
            ax.set_xlim(0.42, 0.62)
            ax.set_title(LABEL[key], fontsize=8)
            ax.set_xlabel("ROC-AUC trên test", fontsize=7.2)
            style_axes(ax)
        axs[0].set_yticks(range(len(kinds)), [KIND[k].replace("Cơ sở: lặp lại hướng phiên trước", "Lặp lại hướng\nphiên trước") for k in kinds], fontsize=7)
        axs[0].invert_yaxis()
        fig.tight_layout()
        savefig(fig, f"ch{chapter}_auc_ci.png")

    # 8) Đường vốn của backtest minh họa
    fig, axs = plt.subplots(1, 3, figsize=(6.8, 2.4))
    series = [("buy_hold", "Mua và giữ", INK, "-"), ("lstm", "LSTM", BLUE, "-"), ("gru", "GRU", AQUA, "--"),
              ("cnn4", "CNN4", ORANGE, ":")]
    for ax, key in zip(axs, MARKETS):
        eq = res.analysis["equity"][key]
        dates = pd.to_datetime(eq["dates"])
        for name, label, color, ls in series:
            ax.plot(dates, eq[name], color=color, ls=ls, lw=1.3, label=label)
        ax.axhline(1, color=MUTED, lw=0.6)
        ax.set_title(LABEL[key], fontsize=8)
        ax.tick_params(axis="x", labelrotation=0)
        ax.xaxis.set_major_locator(matplotlib.dates.YearLocator())
        ax.xaxis.set_major_formatter(matplotlib.dates.DateFormatter("%Y"))
        style_axes(ax)
    axs[0].set_ylabel("Giá trị danh mục (bắt đầu = 1)", fontsize=7.2)
    axs[0].legend(loc="upper left", fontsize=6.6)
    fig.tight_layout()
    savefig(fig, "ch4_equity.png")


# =============================================================================================
# Công thức hiển thị
# =============================================================================================
def render_equation(latex: str, folder: Path = EQUATIONS) -> Path:
    from matplotlib.mathtext import math_to_image
    from matplotlib.font_manager import FontProperties

    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"eq_{hashlib.sha256(latex.encode()).hexdigest()[:12]}.png"
    if not target.exists():
        math_to_image(f"${latex}$", str(target), prop=FontProperties(family="DejaVu Serif", size=13), dpi=300, format="png")
    return target


# =============================================================================================
# Ghi Word
# =============================================================================================
def set_font(target, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    font = target.font
    font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)
    rpr = target.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def add_field(paragraph, instruction, placeholder=""):
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        paragraph._p.append(r)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_with(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run_with(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    text = OxmlElement("w:t")
    text.text = placeholder
    run_with(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def page_number_format(section, fmt, start=1):
    sectpr = section._sectPr
    pg = sectpr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectpr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


def footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, "PAGE", "1")
    for r in p.runs:
        set_font(r, size=11)


class Writer:
    def __init__(self, references: dict):
        self.doc = Document()
        self.references = references
        self.cited: list[str] = []
        self._setup()

    # ---------------------------------------------------------------- styles & page
    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
        styles = self.doc.styles
        normal = styles["Normal"]
        set_font(normal, size=13)
        pf = normal.paragraph_format
        pf.line_spacing, pf.space_after, pf.space_before = 1.3, Pt(6), Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(1.0)
        pf.widow_control = True
        for name, size, align, before, after, italic in [("Heading 1", 15, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, False),
                                                         ("Heading 2", 13.5, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, False),
                                                         ("Heading 3", 13, WD_ALIGN_PARAGRAPH.LEFT, 8, 4, True)]:
            st = styles[name]
            set_font(st, size=size, bold=True, italic=italic, color="000000")
            st.paragraph_format.alignment = align
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.keep_with_next = True
        for name, italic in (("CaptionFigure", True), ("CaptionTable", False)):
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
            set_font(st, size=12, italic=italic)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.line_spacing = 1.1
            st.paragraph_format.space_before = Pt(3) if name == "CaptionFigure" else Pt(8)
            st.paragraph_format.space_after = Pt(10) if name == "CaptionFigure" else Pt(4)
            if name == "CaptionTable":
                st.paragraph_format.keep_with_next = True
        code = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        set_font(code, name="Consolas", size=9)
        code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        code.paragraph_format.first_line_indent = Cm(0)
        code.paragraph_format.left_indent = Cm(0.2)
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.space_before, code.paragraph_format.space_after = Pt(2), Pt(8)
        front = styles.add_style("FrontTitle", WD_STYLE_TYPE.PARAGRAPH)
        front.base_style = normal
        set_font(front, size=14, bold=True)
        front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        front.paragraph_format.first_line_indent = Cm(0)
        front.paragraph_format.space_after = Pt(12)
        lb = styles["List Bullet"]
        set_font(lb, size=13)
        lb.paragraph_format.first_line_indent = Cm(-0.5)
        lb.paragraph_format.left_indent = Cm(1.0)
        lb.paragraph_format.line_spacing = 1.3
        lb.paragraph_format.space_after = Pt(3)
        for toc_level in (1, 2, 3):
            try:
                st = styles[f"TOC {toc_level}"]
            except KeyError:
                st = styles.add_style(f"TOC {toc_level}", WD_STYLE_TYPE.PARAGRAPH)
            set_font(st, size=13, bold=(toc_level == 1))
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.left_indent = Cm(0.6 * (toc_level - 1))
            st.paragraph_format.space_after = Pt(2)
            st.paragraph_format.line_spacing = 1.15
        self.doc.core_properties.title = "Trí tuệ nhân tạo trong đầu tư tài chính — Mở đầu đến Chương 4"
        self.doc.core_properties.author = "Nguyễn Thành Trung"
        self.doc.core_properties.subject = "Tiểu luận môn Phát triển Hệ thống Thông minh"
        self.doc.core_properties.keywords = "AI, học máy, CNN, RNN, LSTM, tài chính, đầu tư, VN-Index"

    # ---------------------------------------------------------------- inline text
    def cite(self, keys: list[str]) -> str:
        numbers = []
        for key in keys:
            key = key.strip().lstrip("@")
            if key not in self.references:
                raise KeyError(f"Thiếu tài liệu tham khảo: {key}")
            if key not in self.cited:
                self.cited.append(key)
            numbers.append(self.cited.index(key) + 1)
        numbers = sorted(set(numbers))
        parts, start = [], None
        for i, n in enumerate(numbers):           # gộp dãy liên tiếp thành [3]–[5]
            if start is None:
                start = n
            if i == len(numbers) - 1 or numbers[i + 1] != n + 1:
                parts.append(f"[{start}]" if start == n else (f"[{start}], [{n}]" if n == start + 1 else f"[{start}]–[{n}]"))
                start = None
        return ", ".join(parts)

    def resolve_citations(self, text: str) -> str:
        return re.sub(r"\[(@[^\]]+)\]", lambda m: self.cite(re.split(r"[;,]\s*", m.group(1))), text)

    def inline(self, paragraph, text, size=None, italic=None):
        text = self.resolve_citations(text)
        pattern = r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*\s][^*]*?\*|`[^`]+`)"
        for piece in re.split(pattern, text):
            if not piece:
                continue
            if piece.startswith("[") and "](" in piece:
                m = re.match(r"\[([^\]]+)\]\((.+)\)", piece)
                self.hyperlink(paragraph, m[1], m[2], size)
                continue
            if piece.startswith("**") and piece.endswith("**"):
                run = paragraph.add_run(piece[2:-2])
                run.bold = True
            elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
                run = paragraph.add_run(piece[1:-1])
                run.italic = True
            elif piece.startswith("`") and piece.endswith("`"):
                run = paragraph.add_run(piece[1:-1])
                set_font(run, name="Consolas", size=(size or 13) - 2.5)
                continue
            else:
                run = paragraph.add_run(piece)
            if size:
                run.font.size = Pt(size)
            if italic:
                run.italic = True

    def hyperlink(self, paragraph, label, url, size=None):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1F4E79")
        props.append(color)
        if size:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(size * 2)))
            props.append(sz)
        run.append(props)
        t = OxmlElement("w:t")
        t.text = label
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    # ---------------------------------------------------------------- blocks
    def paragraph(self, text, style=None, align=None, indent=True):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        self.inline(p, text)
        return p

    def heading(self, text, level, new_page=False):
        p = self.doc.add_heading(level=level)
        p.paragraph_format.page_break_before = new_page
        run = p.add_run(text)
        set_font(run, size={1: 15, 2: 13.5, 3: 13}[level], bold=True, italic=(level == 3), color="000000")
        return p

    def picture(self, path: Path, caption: str, width_cm: float = 15.5):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption(caption, "CaptionFigure")

    def caption(self, text, style):
        p = self.doc.add_paragraph(style=style)
        m = re.match(r"^((?:Hình|Bảng)\s+[\dA-Z]+\.\d+\.)\s*(.*)$", text)
        if m:
            label = p.add_run(m[1] + " ")
            label.bold = True
            label.italic = False
            self.inline(p, m[2])
        else:
            self.inline(p, text)
        return p

    def equation(self, latex: str, number: str | None):
        path = render_equation(latex)
        from PIL import Image

        with Image.open(path) as im:
            width_in, height_in = im.width / 300, im.height / 300
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(6), 1.0
        pf.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        # Giới hạn bề rộng để công thức (căn giữa tại 8 cm) không đè lên số thứ tự căn phải tại 16 cm
        p.add_run().add_picture(str(path), width=Inches(min(width_in, 5.2)))
        if number:
            p.add_run(f"\t({number})")

    def code(self, lines):
        p = self.doc.add_paragraph(style="CodeBlock")
        shade(p._p.get_or_add_pPr(), "F3F3F1")
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()
            run = p.add_run(line)
            set_font(run, name="Consolas", size=9)

    def callout(self, lines):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        tcpr = cell._tc.get_or_add_tcPr()
        shade(tcpr, "EEF4FC")
        borders = OxmlElement("w:tcBorders")
        for edge, val, size, color in (("left", "single", "24", "2A78D6"), ("top", "nil", "0", "auto"),
                                       ("bottom", "nil", "0", "auto"), ("right", "nil", "0", "auto")):
            b = OxmlElement(f"w:{edge}")
            b.set(qn("w:val"), val)
            b.set(qn("w:sz"), size)
            b.set(qn("w:color"), color)
            borders.append(b)
        tcpr.append(borders)
        cell.width = Cm(16)
        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.inline(p, line, size=12)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 0.6

    def table(self, rows, caption=None):
        if caption:
            self.caption(caption, "CaptionTable")
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        lengths = np.array([[len(re.sub(r"\[@[^\]]+\]|\*", "", c)) for c in (r + [""] * (ncols - len(r)))] for r in rows])
        weights = np.clip(np.sqrt(lengths.max(axis=0) + 4), 2.2, None)
        widths = weights / weights.sum() * 16.0
        numeric = re.compile(r"^[−\-+]?[\d.,]+%?( ± [\d.,]+%?)?( \[[−\d.,; ]+\])?$|^—$|^[\d/]+$")
        for i, row in enumerate(rows):
            for j in range(ncols):
                value = row[j] if j < len(row) else ""
                cell = table.cell(i, j)
                cell.width = Cm(widths[j])
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(1), Pt(1)
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or (j > 0 and numeric.match(value.strip())))
                               else WD_ALIGN_PARAGRAPH.LEFT)
                self.inline(p, value, size=10.5)
                if i == 0:
                    for run in p.runs:
                        run.bold = True
                    shade(cell._tc.get_or_add_tcPr(), "DCE6F2")
            tr_pr = table.rows[i]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 0.8

    # ---------------------------------------------------------------- markdown body
    def markdown(self, text, first_heading_new_page=True):
        lines = text.splitlines()
        i, first_h1 = 0, True
        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            if not s or s.startswith("<!--"):
                i += 1
                continue
            if s.startswith("```"):
                i += 1
                block = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i].rstrip("\n"))
                    i += 1
                self.code(block)
                i += 1
                continue
            m = re.match(r"^\$\$(.+)\$\$\s*(?:\((\d+\.\d+)\))?\s*$", s)
            if m:
                self.equation(m[1].strip(), m[2])
                i += 1
                continue
            m = re.match(r"^!\[(.+?)\]\((.+?)\)(?:\{w=([\d.]+)\})?\s*$", s)
            if m:
                path = (DOCS / m[2]).resolve()
                if not path.exists():
                    raise FileNotFoundError(path)
                self.picture(path, m[1], float(m[3]) if m[3] else 15.5)
                i += 1
                continue
            if re.match(r"^Bảng\s+[\dA-Z]+\.\d+\.", s) and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                caption = s
                i += 1
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows, caption)
                continue
            if s.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows)
                continue
            if s.startswith(">"):
                block = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    block.append(lines[i].strip()[1:].strip())
                    i += 1
                self.callout([b for b in block if b])
                continue
            m = re.match(r"^(#{1,3})\s+(.*)$", s)
            if m:
                level = len(m[1])
                new_page = level == 1 and (not first_h1 or first_heading_new_page)
                if level == 1:
                    first_h1 = False
                self.heading(m[2], level, new_page=new_page)
                i += 1
                continue
            if re.match(r"^[-*]\s+", s):
                p = self.doc.add_paragraph(style="List Bullet")
                self.inline(p, re.sub(r"^[-*]\s+", "", s))
                i += 1
                continue
            m = re.match(r"^(\d+)\.\s+(.*)$", s)
            if m:
                p = self.doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.left_indent = Cm(1.1)
                p.paragraph_format.space_after = Pt(3)
                self.inline(p, f"{m[1]}. {m[2]}")
                i += 1
                continue
            block = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||```|!\[|[-*]\s|>|\$\$|\d+\.\s|Bảng\s+\d)", lines[i].strip()):
                block.append(lines[i].strip())
                i += 1
            self.paragraph(" ".join(block))

    # ---------------------------------------------------------------- front & back matter
    def cover(self):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True

        def centered(text, size, bold=False, italic=False, before=0, after=0, caps=False):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            set_font(run, size=size, bold=bold, italic=italic)
            run.font.all_caps = caps
            return p

        centered("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 14, bold=True)
        centered("KHOA CÔNG NGHỆ THÔNG TIN 1", 14, bold=True, after=6)
        centered("———————o0o———————", 12, after=40)
        centered("TIỂU LUẬN MÔN HỌC", 22, bold=True, after=6)
        centered("PHÁT TRIỂN HỆ THỐNG THÔNG MINH", 16, bold=True, after=48)
        centered("ĐỀ TÀI", 13, bold=True, after=6)
        centered("TRÍ TUỆ NHÂN TẠO TRONG ĐẦU TƯ TÀI CHÍNH:", 17, bold=True)
        centered("TỪ HỌC MÁY CƠ BẢN ĐẾN MẠNG TÍCH CHẬP VÀ MẠNG HỒI QUY", 17, bold=True, after=8)
        centered("Phần 1: Mở đầu – Chương 4", 13, italic=True, after=60)
        info = self.doc.add_table(rows=6, cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        data = [("Giảng viên hướng dẫn:", "PGS.TS. Trần Đình Quế"), ("Sinh viên thực hiện:", "Nguyễn Thành Trung"),
                ("Mã sinh viên:", "B23DCCN861"), ("Lớp:", "D23CTPM01-B"), ("Nhóm lớp:", "NhomLop"),
                ("Nhóm tiểu luận:", "NhomTL")]
        for row, (k, val) in zip(info.rows, data):
            for cell, text, bold in ((row.cells[0], k, True), (row.cells[1], val, False)):
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                set_font(run, size=13, bold=bold)
                if text in ("NhomLop", "NhomTL"):
                    run.font.highlight_color = 7  # vàng: nhắc điền trước khi nộp
            row.cells[0].width, row.cells[1].width = Cm(5.5), Cm(7.5)
        centered("Hà Nội, tháng 10 năm 2026", 13, italic=True, before=70)

    def front_matter(self, abbreviations):
        for title, instruction in (("MỤC LỤC", 'TOC \\o "1-3" \\h \\z \\u'),
                                   ("DANH MỤC HÌNH VẼ", 'TOC \\h \\z \\t "CaptionFigure,1"'),
                                   ("DANH MỤC BẢNG BIỂU", 'TOC \\h \\z \\t "CaptionTable,1"')):
            p = self.doc.add_paragraph(title, style="FrontTitle")
            p.paragraph_format.page_break_before = True
            field = self.doc.add_paragraph()
            field.paragraph_format.first_line_indent = Cm(0)
            add_field(field, instruction, "Mở file bằng Microsoft Word và nhấn F9 để cập nhật.")
        p = self.doc.add_paragraph("DANH MỤC TỪ VIẾT TẮT", style="FrontTitle")
        p.paragraph_format.page_break_before = True
        self.table([["Viết tắt", "Tiếng Anh", "Nghĩa tiếng Việt"]] + abbreviations)

    def references_section(self, following_text=""):
        # Phụ lục nằm sau danh mục nhưng nguồn của nó vẫn cần có trong danh mục.
        self.resolve_citations(following_text)
        self.heading("TÀI LIỆU THAM KHẢO", 1, new_page=True)
        for n, key in enumerate(self.cited, 1):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Cm(-1.0)
            pf.left_indent = Cm(1.0)
            pf.line_spacing = 1.15
            pf.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(f"[{n}]\t").font.size = Pt(12)
            pf.tab_stops.add_tab_stop(Cm(1.0))
            self.inline(p, self.references[key], size=12)


ABBREVIATIONS = [
    ["AI", "Artificial Intelligence", "Trí tuệ nhân tạo"], ["ML", "Machine Learning", "Học máy"],
    ["DL", "Deep Learning", "Học sâu"], ["MLP", "Multilayer Perceptron", "Mạng nơ-ron nhiều lớp"],
    ["CNN", "Convolutional Neural Network", "Mạng nơ-ron tích chập"], ["RNN", "Recurrent Neural Network", "Mạng nơ-ron hồi quy"],
    ["LSTM", "Long Short-Term Memory", "Bộ nhớ dài–ngắn hạn"], ["GRU", "Gated Recurrent Unit", "Đơn vị hồi quy có cổng"],
    ["BPTT", "Backpropagation Through Time", "Lan truyền ngược theo thời gian"],
    ["GAF / GASF", "Gramian Angular (Summation) Field", "Trường góc Gram (dạng tổng)"],
    ["BCE", "Binary Cross-Entropy", "Entropy chéo nhị phân"],
    ["ROC-AUC", "Area Under the ROC Curve", "Diện tích dưới đường cong ROC"], ["AP", "Average Precision", "Độ chính xác trung bình"],
    ["CI", "Confidence Interval", "Khoảng tin cậy"], ["RF", "Random Forest", "Rừng ngẫu nhiên"],
    ["OHLC", "Open–High–Low–Close", "Giá mở cửa – cao nhất – thấp nhất – đóng cửa"],
    ["ETF", "Exchange-Traded Fund", "Quỹ hoán đổi danh mục"], ["EMH", "Efficient Market Hypothesis", "Giả thuyết thị trường hiệu quả"],
    ["LLM", "Large Language Model", "Mô hình ngôn ngữ lớn"], ["SSM", "State Space Model", "Mô hình không gian trạng thái"],
    ["UCI", "UC Irvine Machine Learning Repository", "Kho dữ liệu học máy của Đại học California, Irvine"],
    ["HOSE", "Ho Chi Minh City Stock Exchange", "Sở Giao dịch Chứng khoán TP. Hồ Chí Minh"],
]


def substitute(text: str, values: dict, tables: dict) -> str:
    def value(m):
        key = m.group(1)
        if key not in values:
            raise KeyError(f"Thiếu số liệu {{{{v:{key}}}}}")
        return values[key]

    def table(m):
        key, number = m.group(1), m.group(2)
        if key not in tables:
            raise KeyError(f"Thiếu bảng {{{{t:{key}}}}}")
        return tables[key].replace("{n}", number)

    text = re.sub(r"\{\{t:([\w.]+):([\dA-Z]+\.\d+)\}\}", table, text)
    text = re.sub(r"\{\{v:([\w.\-]+)\}\}", value, text)
    leftover = re.findall(r".{0,40}[{][{][tv]:.{0,40}", text)   # công thức LaTeX có "}}" hợp lệ nên chỉ bắt {{t: và {{v:
    if leftover:
        raise ValueError(f"Còn chỗ giữ chỗ chưa thay: {leftover[:3]}")
    return text


def build():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    res = Results()
    make_figures(res)
    values, tables = compute_values(res), compute_tables(res, compute_values(res))
    references = json.loads((DOCS / "references.json").read_text(encoding="utf8"))
    files = sorted(CHAPTERS.glob("*.md"))
    texts = [substitute(p.read_text(encoding="utf8"), values, tables) for p in files]
    w = Writer(references)
    w.cover()
    w.front_matter(ABBREVIATIONS)
    main = w.doc.add_section(WD_SECTION.NEW_PAGE)
    footer_page_number(w.doc.sections[0])
    page_number_format(w.doc.sections[0], "lowerRoman", 1)
    footer_page_number(main)
    page_number_format(main, "decimal", 1)
    appendix_text = None
    for path, text in zip(files, texts):
        if "phu_luc" in path.name:
            appendix_text = text
            continue
        w.markdown(text, first_heading_new_page=(path != files[0]))
    w.references_section(appendix_text or "")
    if appendix_text:
        w.markdown(appendix_text)
    out = DOCS / f"{NAME}.docx"
    w.doc.save(out)
    merged = "\n\n".join(texts)
    (DOCS / f"{NAME}.md").write_text(merged, encoding="utf8")
    unused = sorted(set(references) - set(w.cited))
    print(json.dumps({"docx": str(out), "words": len(merged.split()), "references_cited": len(w.cited),
                      "references_unused": unused}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    build()
